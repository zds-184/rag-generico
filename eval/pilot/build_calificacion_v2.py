"""Genera la plantilla de calificación v2 (protocolo v2) a partir de una corrida existente.

Lee (solo lectura) questions.jsonl, <corrida>/traces.jsonl, <corrida>/config.json, evidence_pilot.json, data/manifest.json
y, si existe, <corrida>/n1_v2.json. Escribe un libro NUEVO en eval/pilot/v2/ y se NIEGA a sobrescribir uno existente
(las calificaciones manuales no se pierden). No toca nada de la corrida original.

Hojas: Guia, Respuestas (una fila por pregunta x modo; columnas calculadas), Afirmaciones (donde se califica a mano),
Hechos (referencia descompuesta + `llegó`), Contextos, Resumen, Config.

Uso:
  python eval/pilot/build_calificacion_v2.py eval/pilot/results/20260922_064359
  python eval/pilot/build_calificacion_v2.py eval/pilot/results/20260922_064359 --demo   # relleno con los 27 ejemplos
"""
import csv
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
ROOT = HERE.parents[1]
OUT_DIR = HERE / "v2"
NAFF = 400          # filas disponibles en Afirmaciones
NHEC = 200          # filas máximas leídas en Hechos

HDR = PatternFill("solid", fgColor="1F3864")
GREEN = PatternFill("solid", fgColor="E2EFDA")     # automático
YELLOW = PatternFill("solid", fgColor="FFF2CC")    # manual
GREY = PatternFill("solid", fgColor="F2F2F2")
WHITE_B = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F10 = Font(name="Arial", size=10)
F10B = Font(name="Arial", size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
UNLOCK = Protection(locked=False)


def head(ws, names, widths):
    for i, (n, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(1, i, n)
        c.font, c.fill, c.alignment = WHITE_B, HDR, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"


def col(names):
    return {n: get_column_letter(i) for i, n in enumerate(names, 1)}


def main(run, demo):
    run = Path(run)
    name = "calificacion_v2_DEMO.xlsx" if demo else "calificacion_v2_piloto.xlsx"
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / name
    if out.exists():
        sys.exit(f"{out} ya existe: no se sobrescribe (protege las calificaciones). Renómbralo o muévelo si quieres regenerar.")

    questions = {}
    for line in (HERE / "questions.jsonl").read_text(encoding="utf8").splitlines():
        q = json.loads(line)
        questions[q["id"]] = q
    traces = [json.loads(line) for line in (run / "traces.jsonl").read_text(encoding="utf8").splitlines()]
    cfg = json.loads((run / "config.json").read_text(encoding="utf8"))
    ev = json.loads((HERE / "evidence_pilot.json").read_text(encoding="utf8"))
    manifest = json.loads((ROOT / "data" / "manifest.json").read_text(encoding="utf8"))
    n1 = json.loads((run / "n1_v2.json").read_text(encoding="utf8")) if (run / "n1_v2.json").exists() else {}

    wb = Workbook()
    wsG = wb.active
    wsG.title = "Guia"
    wsR = wb.create_sheet("Respuestas")
    wsA = wb.create_sheet("Afirmaciones")
    wsH = wb.create_sheet("Hechos")
    wsC = wb.create_sheet("Contextos")
    wsS = wb.create_sheet("Resumen")
    wsK = wb.create_sheet("Config")

    # ------------------------------------------------------------------ Afirmaciones
    A_names = ["qid", "modo", "n", "afirmación", "tipo", "vs_ref", "hecho_ref", "zona", "respaldo", "cita_n",
               "fragmento_sostiene", "nota", "códigos (auto)", "chequeo (auto)"]
    A = col(A_names)
    head(wsA, A_names, [7, 7, 5, 60, 13, 14, 14, 10, 17, 9, 13, 40, 18, 28])
    rngA = lambda k: f"Afirmaciones!${A[k]}$2:${A[k]}${NAFF + 1}"
    for r in range(2, NAFF + 2):
        wsA[f"{A['n']}{r}"] = f'=IF({A["qid"]}{r}="","",COUNTIFS(${A["qid"]}$2:{A["qid"]}{r},{A["qid"]}{r},${A["modo"]}$2:{A["modo"]}{r},{A["modo"]}{r}))'
        X = (f'IF(AND({A["tipo"]}{r}="hecho",{A["vs_ref"]}{r}="INCORRECTA"),"INC ","")'
             f'&IF(AND({A["tipo"]}{r}="abstencion",{A["vs_ref"]}{r}="INCORRECTA"),"ABS-E ","")'
             f'&IF(AND({A["tipo"]}{r}="hecho",{A["vs_ref"]}{r}="FUERA_REF",{A["zona"]}{r}="no_resp"),"NO-ABS ","")'
             f'&IF(AND({A["tipo"]}{r}="hecho",{A["respaldo"]}{r}="SIN-RESPALDO"),"SIN-R ","")'
             f'&IF(AND({A["tipo"]}{r}="hecho",{A["respaldo"]}{r}="CITA-X"),"CITA-X ","")'
             f'&IF(AND({A["tipo"]}{r}="hecho",{A["respaldo"]}{r}="SIN-CITA"),"SIN-C ","")')
        wsA[f"{A['códigos (auto)']}{r}"] = (
            f'=IF({A["afirmación"]}{r}="","",IF(OR({A["tipo"]}{r}="",{A["vs_ref"]}{r}=""),"",'
            f'IF(TRIM({X})="","-",TRIM({X}))))')
        wsA[f"{A['chequeo (auto)']}{r}"] = (
            f'=IF({A["afirmación"]}{r}="","",IF(OR({A["qid"]}{r}="",{A["modo"]}{r}="",{A["tipo"]}{r}="",{A["vs_ref"]}{r}=""),"falta qid, modo, tipo o vs_ref",'
            f'IF(AND({A["modo"]}{r}="llm",{A["respaldo"]}{r}<>""),"respaldo no aplica en modo LLM",'
            f'IF(AND({A["tipo"]}{r}="abstencion",{A["respaldo"]}{r}<>""),"una abstención no se califica en citas",'
            f'IF(AND({A["modo"]}{r}="rag",{A["tipo"]}{r}="hecho",{A["respaldo"]}{r}=""),"falta respaldo","ok")))))')
        for k in A_names:
            c = wsA[f"{A[k]}{r}"]
            c.font, c.alignment = F10, WRAP
            if k in ("n", "códigos (auto)", "chequeo (auto)"):
                c.fill = GREEN
            else:
                c.fill = YELLOW
                c.protection = UNLOCK
    for key, lst in (("qid", ",".join(sorted(questions))), ("modo", "rag,llm"), ("tipo", "hecho,abstencion"),
                     ("vs_ref", "CORRECTA,INCORRECTA,FUERA_REF"), ("zona", "ref,no_resp"),
                     ("respaldo", "SOSTENIDA,CITA-X,SIN-CITA,SIN-RESPALDO")):
        dv = DataValidation(type="list", formula1=f'"{lst}"', allow_blank=True, showErrorMessage=True)
        wsA.add_data_validation(dv)
        dv.add(f"{A[key]}2:{A[key]}{NAFF + 1}")

    # ------------------------------------------------------------------ Hechos
    H_names = ["qid", "hecho", "texto del hecho", "evidencia (n frag.)", "llegó (RAG)", "cubierto RAG", "cubierto LLM",
               "cubierto y citado RAG"]
    Hc = col(H_names)
    head(wsH, H_names, [7, 7, 62, 12, 12, 12, 12, 14])
    rngH = lambda k: f"Hechos!${Hc[k]}$2:${Hc[k]}${NHEC + 1}"
    hrow = 2
    for qid in sorted(ev):
        arrived = {h["fact"]: h["llego"] for h in n1.get(qid, {}).get("hechos", [])}
        for f in ev[qid]:
            wsH.cell(hrow, 1, qid)
            wsH.cell(hrow, 2, f["fact"])
            wsH.cell(hrow, 3, f["text"])
            wsH.cell(hrow, 4, len(f["evidence"]))
            wsH.cell(hrow, 5, arrived.get(f["fact"], ""))
            fr = lambda k, hr=hrow: f"${Hc[k]}{hr}"
            hr_ = f'SUBSTITUTE({rngA("hecho_ref")}," ","")'
            base = (f'({rngA("qid")}={fr("qid")})*({rngA("vs_ref")}="CORRECTA")*({rngA("zona")}<>"no_resp")'
                    f'*ISNUMBER(SEARCH(";"&{fr("hecho")}&";",";"&{hr_}&";"))')
            wsH[f"{Hc['cubierto RAG']}{hrow}"] = (
                f'=IF(COUNTIFS({rngA("qid")},{fr("qid")},{rngA("modo")},"rag")=0,"",'
                f'IF(SUMPRODUCT({base}*({rngA("modo")}="rag"))>0,1,0))')
            wsH[f"{Hc['cubierto LLM']}{hrow}"] = (
                f'=IF(COUNTIFS({rngA("qid")},{fr("qid")},{rngA("modo")},"llm")=0,"",'
                f'IF(SUMPRODUCT({base}*({rngA("modo")}="llm"))>0,1,0))')
            wsH[f"{Hc['cubierto y citado RAG']}{hrow}"] = (
                f'=IF(COUNTIFS({rngA("qid")},{fr("qid")},{rngA("modo")},"rag")=0,"",'
                f'IF(SUMPRODUCT({base}*({rngA("modo")}="rag")*({rngA("respaldo")}="SOSTENIDA"))>0,1,0))')
            for k in H_names:
                c = wsH[f"{Hc[k]}{hrow}"]
                c.font, c.alignment = F10, WRAP
                c.fill = YELLOW if k == "llegó (RAG)" else (GREEN if k.startswith("cubierto") else GREY)
                if k == "llegó (RAG)":
                    c.protection = UNLOCK
            hrow += 1
    dv = DataValidation(type="list", formula1='"SI,ALT,ALT-P,NO"', allow_blank=True)
    wsH.add_data_validation(dv)
    dv.add(f"{Hc['llegó (RAG)']}2:{Hc['llegó (RAG)']}{NHEC + 1}")

    # ------------------------------------------------------------------ Respuestas
    R_names = ["qid", "modo", "idioma", "N0 corpus", "pregunta", "respuesta de referencia", "parte NO respondible",
               "RESPUESTA DEL SISTEMA", "truncada", "N1 (auto)", "hechos llegaron (auto)",
               "n afirm.", "hechos H", "cubiertos C", "cobertura", "INC", "ABS-E", "NO-ABS", "FUERA_REF",
               "SIN-R", "CITA-X", "SIN-C", "afirm. de hecho", "de hecho sostenidas",
               "N2", "N3", "alertas", "diagnóstico (lista)", "primer eslabón",
               "precisión citas", "cobertura citas", "aux E1", "aux E2", "comentario"]
    Rc = col(R_names)
    head(wsR, R_names, [6, 6, 6, 9, 38, 38, 26, 60, 8, 10, 11, 7, 7, 8, 9, 6, 6, 7, 8, 6, 7, 6, 8, 9, 5, 5, 22, 18, 10, 9, 9, 6, 6, 36])
    rngR = lambda k: f"Respuestas!${Rc[k]}$2:${Rc[k]}$200"
    order = []
    for qid in sorted(questions):
        for mode in ("rag", "llm"):
            t = next((t for t in traces if t["eval"]["id"] == qid and t["mode"] == mode), None)
            if t:
                order.append((qid, mode, t))
    nr = len(order)
    r = 2
    for qid, mode, t in order:
        q = questions[qid]
        trunc = "SI" if t["usage"]["completion_tokens"] >= t.get("params", {}).get("num_predict", 1024) else "NO"
        vals = {"qid": qid, "modo": mode, "idioma": q["lang"].upper(),
                "N0 corpus": {"complete": "completa", "partial": "parcial", "none": "nula"}.get(q["answerability"], q["answerability"]),
                "pregunta": q["question"], "respuesta de referencia": q["reference"],
                "parte NO respondible": q["unanswerable_part"] or "", "RESPUESTA DEL SISTEMA": t["answer"],
                "truncada": trunc}
        if mode == "rag" and qid in n1:
            vals["N1 (auto)"] = n1[qid]["n1"]
            vals["hechos llegaron (auto)"] = n1[qid]["hechos_llegaron"]
        for k, v in vals.items():
            wsR[f"{Rc[k]}{r}"] = v
        c = lambda k, rr=r: f"${Rc[k]}{rr}"
        cnt = lambda *conds: "COUNTIFS(" + f'{rngA("qid")},{c("qid")},{rngA("modo")},{c("modo")},' + ",".join(conds) + ")"
        T = lambda k, v: f'{rngA(k)},"{v}"'
        f = {
            "n afirm.": f'=COUNTIFS({rngA("qid")},{c("qid")},{rngA("modo")},{c("modo")})',
            "hechos H": f'=COUNTIFS({rngH("qid")},{c("qid")})',
            "cubiertos C": f'=IF({c("modo")}="rag",SUMIFS({rngH("cubierto RAG")},{rngH("qid")},{c("qid")}),SUMIFS({rngH("cubierto LLM")},{rngH("qid")},{c("qid")}))',
            "cobertura": f'=IF({c("hechos H")}=0,"-",{c("cubiertos C")}&"/"&{c("hechos H")})',
            "INC": "=" + cnt(T("tipo", "hecho"), T("vs_ref", "INCORRECTA")),
            "ABS-E": "=" + cnt(T("tipo", "abstencion"), T("vs_ref", "INCORRECTA")),
            "NO-ABS": "=" + cnt(T("tipo", "hecho"), T("vs_ref", "FUERA_REF"), T("zona", "no_resp")),
            "FUERA_REF": "=" + cnt(T("tipo", "hecho"), T("vs_ref", "FUERA_REF")),
            "SIN-R": "=" + cnt(T("tipo", "hecho"), T("respaldo", "SIN-RESPALDO")),
            "CITA-X": "=" + cnt(T("tipo", "hecho"), T("respaldo", "CITA-X")),
            "SIN-C": "=" + cnt(T("tipo", "hecho"), T("respaldo", "SIN-CITA")),
            "afirm. de hecho": "=" + cnt(T("tipo", "hecho")),
            "de hecho sostenidas": "=" + cnt(T("tipo", "hecho"), T("respaldo", "SOSTENIDA")),
            "N2": (f'=IF({c("n afirm.")}=0,"",IF({c("hechos H")}=0,IF({c("NO-ABS")}>0,0,2),'
                   f'IF(AND({c("cubiertos C")}={c("hechos H")},{c("INC")}+{c("ABS-E")}+{c("NO-ABS")}=0),2,'
                   f'IF({c("cubiertos C")}=0,0,1))))'),
            "N3": (f'=IF({c("n afirm.")}=0,"",IF({c("modo")}="llm","N/A",IF({c("afirm. de hecho")}=0,"N/A",'
                   f'IF({c("de hecho sostenidas")}={c("afirm. de hecho")},2,IF({c("de hecho sostenidas")}=0,0,1)))))'),
            "alertas": (f'=IF({c("n afirm.")}=0,"",IF(AND({c("INC")}+{c("ABS-E")}+{c("NO-ABS")}=0,{c("SIN-R")}+{c("CITA-X")}=0,{c("truncada")}<>"SI"),"-",'
                        f'TRIM(IF({c("INC")}+{c("ABS-E")}+{c("NO-ABS")}>0,"CRITICA ","")&IF({c("SIN-R")}+{c("CITA-X")}>0,"RESPALDO ","")&IF({c("truncada")}="SI","TRUNC",""))))'),
            "aux E1": (f'=IF({c("modo")}="rag",SUMPRODUCT(({rngH("qid")}={c("qid")})*({rngH("cubierto RAG")}=0)'
                       f'*(({rngH("llegó (RAG)")}="NO")+({rngH("llegó (RAG)")}="ALT-P")+({rngH("llegó (RAG)")}=""))),0)'),
            "aux E2": (f'=IF({c("modo")}="rag",SUMPRODUCT(({rngH("qid")}={c("qid")})*({rngH("cubierto RAG")}=0)'
                       f'*(({rngH("llegó (RAG)")}="SI")+({rngH("llegó (RAG)")}="ALT"))),{c("hechos H")}-{c("cubiertos C")})'),
            "diagnóstico (lista)": (f'=IF({c("n afirm.")}=0,"",IF(AND({c("NO-ABS")}=0,{c("aux E1")}=0,{c("aux E2")}=0,{c("INC")}+{c("ABS-E")}=0,'
                                    f'{c("SIN-R")}+{c("CITA-X")}+{c("SIN-C")}=0),"OK",'
                                    f'TRIM(IF({c("NO-ABS")}>0,"E0 ","")&IF({c("aux E1")}>0,"E1 ","")'
                                    f'&IF(OR({c("aux E2")}>0,AND({c("INC")}+{c("ABS-E")}>0,{c("aux E1")}=0)),"E2 ","")'
                                    f'&IF({c("SIN-R")}+{c("CITA-X")}+{c("SIN-C")}>0,"E3 ",""))))'),
            "primer eslabón": f'=IF({c("diagnóstico (lista)")}="","",IF({c("diagnóstico (lista)")}="OK","OK",LEFT({c("diagnóstico (lista)")},2)))',
            "precisión citas": f'=IF(OR({c("n afirm.")}=0,{c("modo")}="llm",{c("afirm. de hecho")}=0),"-",{c("de hecho sostenidas")}&"/"&{c("afirm. de hecho")})',
            "cobertura citas": (f'=IF(OR({c("n afirm.")}=0,{c("modo")}<>"rag",{c("hechos H")}=0),"-",'
                                f'SUMIFS({rngH("cubierto y citado RAG")},{rngH("qid")},{c("qid")})&"/"&{c("hechos H")})'),
        }
        for k, v in f.items():
            wsR[f"{Rc[k]}{r}"] = v
        for k in R_names:
            cell = wsR[f"{Rc[k]}{r}"]
            cell.font, cell.alignment = F10, WRAP
            cell.fill = YELLOW if k == "comentario" else (GREEN if k in f or k in ("N1 (auto)", "hechos llegaron (auto)", "truncada") else GREY)
            if k == "comentario":
                cell.protection = UNLOCK
        wsR.row_dimensions[r].height = 150
        r += 1

    # ------------------------------------------------------------------ Contextos
    C_names = ["qid", "rank", "point_id", "attack_id", "chunk", "score", "texto del fragmento"]
    head(wsC, C_names, [7, 6, 38, 11, 8, 8, 110])
    cr = 2
    pid2meta = {}
    for qid in ev:
        for f in ev[qid]:
            for e in f["evidence"]:
                pid2meta[e["point_id"]] = e
    for qid, mode, t in order:
        if mode != "rag":
            continue
        for ctx in t["contexts"]:
            md = ctx.get("metadata") or {}
            for j, v in enumerate([qid, ctx["rank"], ctx["point_id"], md.get("attack_id", ""),
                                   f"{md.get('chunk_index', '')}/{md.get('n_chunks', '')}", round(ctx.get("score", 0), 4), ctx["text"]], 1):
                cell = wsC.cell(cr, j, v)
                cell.font, cell.alignment = F10, WRAP
            wsC.row_dimensions[cr].height = 120
            cr += 1

    # ------------------------------------------------------------------ Resumen
    wsS.column_dimensions["A"].width = 58
    wsS.column_dimensions["B"].width = 16
    wsS.column_dimensions["C"].width = 16
    for j, v in enumerate(["Lectura (protocolo v2 §5)", "RAG", "LLM"], 1):
        c = wsS.cell(1, j, v)
        c.font, c.fill = WHITE_B, HDR
    B, J = rngR("modo"), rngR("n afirm.")
    sm = lambda k, m, extra="": f'SUMIFS({rngR(k)},{B},"{m}",{J},">0"{extra})'
    rows = [
        ("Respuestas calificadas", lambda m: f'=COUNTIFS({B},"{m}",{J},">0")'),
        ("A. Hechos cubiertos ÷ hechos (respondibles)", lambda m: f'=IF({sm("hechos H", m)}=0,"-",{sm("cubiertos C", m)}/{sm("hechos H", m)})'),
        ("A. Afirmaciones INCORRECTA ÷ afirmaciones de hecho", lambda m: f'=IF({sm("afirm. de hecho", m)}=0,"-",{sm("INC", m)}/{sm("afirm. de hecho", m)})'),
        ("B. Respuestas con abstención errónea (ABS-E)", lambda m: f'=COUNTIFS({B},"{m}",{rngR("ABS-E")},">0")'),
        ("B. Respuestas con respuesta sin base (NO-ABS)", lambda m: f'=COUNTIFS({B},"{m}",{rngR("NO-ABS")},">0")'),
        ("C. Precisión de citas (sostenidas ÷ afirmaciones de hecho)", lambda m: '="N/A"' if m == "llm" else f'=IF({sm("afirm. de hecho", m)}=0,"-",{sm("de hecho sostenidas", m)}/{sm("afirm. de hecho", m)})'),
        ("C. Cobertura de citas (hechos cubiertos y citados ÷ hechos)", lambda m: '="N/A"' if m == "llm" else f'=IF({sm("hechos H", m)}=0,"-",SUM({rngH("cubierto y citado RAG")})/{sm("hechos H", m)})'),
        ("C. Afirmaciones SIN-RESPALDO", lambda m: '="N/A"' if m == "llm" else f'=SUMIFS({rngR("SIN-R")},{B},"{m}")'),
        ("C. Afirmaciones CITA-X", lambda m: '="N/A"' if m == "llm" else f'=SUMIFS({rngR("CITA-X")},{B},"{m}")'),
        ("C. Afirmaciones SIN-CITA", lambda m: '="N/A"' if m == "llm" else f'=SUMIFS({rngR("SIN-C")},{B},"{m}")'),
        ("D. Afirmaciones FUERA_REF", lambda m: f'=SUMIFS({rngR("FUERA_REF")},{B},"{m}")'),
        ("N2 promedio (descriptivo, 0-2)", lambda m: f'=IFERROR(AVERAGEIFS({rngR("N2")},{B},"{m}"),"-")'),
        ("N3 promedio (descriptivo, 0-2; excluye N/A)", lambda m: '="N/A"' if m == "llm" else f'=IFERROR(AVERAGEIFS({rngR("N3")},{B},"{m}"),"-")'),
        ("Alerta CRITICA (respuestas)", lambda m: f'=COUNTIFS({B},"{m}",{rngR("alertas")},"*CRITICA*")'),
        ("Alerta RESPALDO (respuestas)", lambda m: f'=COUNTIFS({B},"{m}",{rngR("alertas")},"*RESPALDO*")'),
        ("Alerta TRUNC (respuestas)", lambda m: f'=COUNTIFS({B},"{m}",{rngR("alertas")},"*TRUNC*")'),
    ]
    for i, (label, fn) in enumerate(rows, 2):
        wsS.cell(i, 1, label).font = F10
        for j, m in ((2, "rag"), (3, "llm")):
            c = wsS.cell(i, j, fn(m))
            c.font = F10
            if "÷" in label:
                c.number_format = "0.00"
            if "promedio" in label:
                c.number_format = "0.00"
    wsS.cell(len(rows) + 3, 1, "Las lecturas A-D y las alertas se calculan solas desde la hoja Afirmaciones. Con pocas respuestas son conteos descriptivos, sin pruebas de significancia.").font = F10
    wsS.cell(len(rows) + 3, 1).alignment = WRAP
    wsS.row_dimensions[len(rows) + 3].height = 45

    # ------------------------------------------------------------------ Config
    wsK.column_dimensions["A"].width = 34
    wsK.column_dimensions["B"].width = 80
    info = [
        ("rúbrica", "protocolo v2 (borrador) - eval/pilot/protocol_v2.md"),
        ("plantilla creada", datetime.now().isoformat(timespec="seconds")),
        ("corrida de origen", str(run.name)),
        ("evaluador", ""),
        ("fecha de calificación", ""),
        ("", ""),
        ("CORRIDA (config.json)", ""),
    ] + [(k, str(v)) for k, v in cfg.items()] + [("", ""), ("CORPUS (data/manifest.json)", "")] + [(k, str(v)) for k, v in manifest.items()]
    for i, (k, v) in enumerate(info, 1):
        wsK.cell(i, 1, k).font = F10B if v == "" else F10
        wsK.cell(i, 2, v).font = F10
        if k in ("evaluador", "fecha de calificación"):
            wsK.cell(i, 2).fill = YELLOW
            wsK.cell(i, 2).protection = UNLOCK

    # ------------------------------------------------------------------ Guia
    wsG.column_dimensions["A"].width = 26
    wsG.column_dimensions["B"].width = 120
    guia = [
        ("Para qué sirve", "Calificar cada respuesta POR AFIRMACIÓN (hoja Afirmaciones). N2, N3, códigos, alertas y diagnóstico se calculan solos en Respuestas. Reglas completas: protocol_v2.md §3-§5."),
        ("Colores", "Amarillo = lo escribes o eliges tú. Verde = automático (no editable). Gris = dato de la pregunta/corrida."),
        ("Flujo", "1) En Respuestas, lee pregunta, referencia y respuesta del sistema (en RAG, los fragmentos están en Contextos). 2) En Afirmaciones, una fila por afirmación (máx. ~6): qid, modo, texto, tipo, vs_ref, hecho_ref, zona, respaldo, cita_n, fragmento_sostiene. 3) Revisa la columna 'chequeo'. 4) Revisa 'llegó' en Hechos (sobre todo ALT-P). 5) Lee N2/N3/alertas/diagnóstico en Respuestas."),
        ("tipo", "hecho = afirma algo. abstencion = dice que el contexto o el modelo no tiene la información ('los pasajes no dicen...'). Afirmar que ATT&CK no documenta algo ('ATT&CK no lista...') es un HECHO y puede ser INCORRECTA."),
        ("vs_ref", "CORRECTA = coincide con la referencia (en una abstención: se abstiene donde debía). INCORRECTA = contradice la referencia (en una abstención: niega algo que la referencia sí contiene). FUERA_REF = ni confirma ni contradice la referencia."),
        ("hecho_ref", "H# de la hoja Hechos que cubre una afirmación CORRECTA; varios separados por ';' (ej. H2;H3). En INCORRECTA/abstención es opcional (nota)."),
        ("zona", "ref (por defecto) o no_resp si responde a la parte NO respondible. Un hecho FUERA_REF en no_resp genera NO-ABS."),
        ("respaldo (solo RAG, solo hecho)", "SOSTENIDA = tiene cita [n] y ese fragmento expresa la relación. CITA-X = cita pero ese fragmento no la sostiene. SIN-CITA = sin cita, pero algún fragmento recibido la sostiene. SIN-RESPALDO = ningún fragmento recibido la sostiene (tenga o no cita). Dos nombres en la misma lista no sostienen 'A usa B'."),
        ("llegó (hoja Hechos)", "Prellenado desde evidence_check.py: SI = llegó un fragmento gold; ALT = llegó evidencia alternativa; ALT-P = solo evidencia parcial (cuenta como NO salvo que la cambies a ALT y lo justifiques en comentario); NO = no llegó. Verifica leyendo Contextos."),
        ("Códigos", "INC, ABS-E, NO-ABS, SIN-R, CITA-X, SIN-C se derivan de las columnas. Omisiones = hechos de la referencia sin cubrir (cubierto = 0)."),
        ("N2", "H=0: 2 si no hay NO-ABS, si no 0. H>0: 2 si cubre todos y no hay INC/ABS-E/NO-ABS; 0 si no cubre ninguno; 1 en los demás casos. Sin topes."),
        ("N3", "Solo RAG, sobre afirmaciones de hecho: 2 todas sostenidas; 0 ninguna; 1 parcial; N/A sin afirmaciones de hecho."),
        ("Alertas", "CRITICA = hay INC, ABS-E o NO-ABS. RESPALDO = hay SIN-R o CITA-X. TRUNC = respuesta truncada por el límite de tokens. No cambian N2 ni N3."),
        ("Diagnóstico", "Lista: E0 (NO-ABS) · E1 (hecho sin cubrir cuya evidencia no llegó) · E2 (hecho sin cubrir cuya evidencia sí llegó; o INC/ABS-E con todo lo demás llegado) · E3 (SIN-R, CITA-X o SIN-C). OK si no hay ninguno. Modo LLM: solo E0 y E2."),
        ("Protección", "Las hojas están protegidas SIN contraseña para evitar editar fórmulas por accidente (Revisar > Desproteger hoja). Este archivo nunca se sobrescribe desde el script: guarda copias de seguridad de tus calificaciones."),
        ("Verificación", "Las fórmulas se evaluaron con un motor independiente (librería Python 'formulas') sobre calificacion_v2_DEMO.xlsx, que trae las 27 afirmaciones de los ejemplos: N2, N3, alertas y diagnóstico de las 8 respuestas coinciden con la columna 'esperado' de Respuestas. No se han abierto todavía en Excel; si algo no calcula en tu Excel, avísalo."),
    ]
    for i, (k, v) in enumerate(guia, 1):
        wsG.cell(i, 1, k).font = F10B
        c = wsG.cell(i, 2, v)
        c.font, c.alignment = F10, WRAP
        wsG.cell(i, 1).alignment = WRAP
        wsG.row_dimensions[i].height = 62 if len(v) > 220 else (46 if len(v) > 120 else 30)

    # ------------------------------------------------------------------ demo: rellena con el CSV de ejemplos
    if demo:
        rows_csv = list(csv.DictReader(open(HERE / "afirmaciones_ejemplos.csv", newline="", encoding="utf-8-sig")))
        for i, rw in enumerate(rows_csv, 2):
            for k in ("qid", "modo", "afirmacion", "tipo", "vs_ref", "hecho_ref", "zona", "respaldo", "cita_n", "fragmento_sostiene", "nota"):
                key = {"afirmacion": "afirmación"}.get(k, k)
                wsA[f"{A[key]}{i}"] = rw[k]
        expected = {  # N2, N3, diagnóstico, de los ejemplos calificados (ejemplos_calificados_v2.md)
            ("p01", "rag"): (0, "N/A", "E1"), ("p01", "llm"): (1, "N/A", "E2"),
            ("p05", "rag"): (1, 1, "E1 E2 E3"), ("p05", "llm"): (0, "N/A", "E2"),
            ("p06", "rag"): (1, 2, "E2"), ("p06", "llm"): (0, "N/A", "E2"),
            ("p09", "rag"): (0, 2, "E1 E2"), ("p09", "llm"): (1, "N/A", "E0 E2"),
        }
        ecol = len(R_names) + 1
        c = wsR.cell(1, ecol, "ESPERADO (ejemplos): N2 | N3 | diagnóstico")
        c.font, c.fill = WHITE_B, HDR
        wsR.column_dimensions[get_column_letter(ecol)].width = 26
        for rr, (qid, mode, t) in enumerate(order, 2):
            if (qid, mode) in expected:
                n2, n3, dg = expected[(qid, mode)]
                wsR.cell(rr, ecol, f"{n2} | {n3} | {dg}").font = F10B

    for ws in (wsG, wsR, wsA, wsH, wsC, wsS, wsK):
        ws.protection.sheet = True
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
        ws.protection.sort = False
        ws.protection.autoFilter = False
    wb.save(out)
    print(f"escrito {out}  ({nr} respuestas, {hrow - 2} hechos, {NAFF} filas de afirmaciones)")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    main(args[0], "--demo" in sys.argv)
