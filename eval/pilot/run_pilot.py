"""Ejecuta el piloto de evaluación diagnóstica contra la API del RAG.

Lee eval/pilot/questions.jsonl, pregunta en modo rag y llm con una configuración fija,
guarda la traza completa de cada respuesta (contextos, prompt, tokens, tiempos) y genera
la hoja de calificación manual (Excel) con el nivel de recuperación (N1) prellenado.

Uso:
  python eval/pilot/run_pilot.py                      # corrida nueva en eval/pilot/results/<timestamp>/
  python eval/pilot/run_pilot.py --resume <carpeta>   # continúa una corrida interrumpida
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
ANSWERABILITY_ES = {"complete": "completa", "partial": "parcial", "none": "nula"}
DIAGNOSES = ["OK", "E0-corpus", "E1-recuperacion", "E2-generacion", "E3-citas"]


def log(msg: str) -> None:
    print(msg.encode("ascii", "replace").decode(), flush=True)


def load_questions(path: Path) -> list[dict]:
    return [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]


def load_done(traces_path: Path) -> dict[tuple[str, str], dict]:
    done = {}
    if traces_path.exists():
        for l in traces_path.read_text(encoding="utf-8").splitlines():
            if l.strip():
                r = json.loads(l)
                done[(r["eval"]["id"], r["mode"])] = r
    return done


def auto_retrieval_eval(q: dict, contexts: list[dict]) -> dict:
    gold_points = set(q.get("gold_point_ids") or [])
    gold_ids = set(q.get("gold_attack_ids") or [])
    point_ranks = [c["rank"] for c in contexts if c["point_id"] in gold_points]
    attack_ranks = [c["rank"] for c in contexts if c["metadata"].get("attack_id") in gold_ids]
    return {
        "gold_point_hit": bool(point_ranks),
        "gold_point_first_rank": min(point_ranks) if point_ranks else None,
        "gold_points_found": len(point_ranks),
        "gold_points_total": len(gold_points),
        "gold_attack_hit": bool(attack_ranks),
        "gold_attack_first_rank": min(attack_ranks) if attack_ranks else None,
        "retrieved_attack_ids": [c["metadata"].get("attack_id") for c in contexts],
    }


def run(args: argparse.Namespace) -> Path:
    questions = load_questions(Path(args.questions))
    run_dir = Path(args.resume) if args.resume else HERE / "results" / f"{datetime.now():%Y%m%d_%H%M%S}"
    run_dir.mkdir(parents=True, exist_ok=True)
    traces_path = run_dir / "traces.jsonl"
    done = load_done(traces_path)

    health = requests.get(f"{args.api}/health", timeout=30).json()
    config = {
        "api": args.api,
        "model": args.model or health["defaults"]["llm_model"],
        "embed_model": health["defaults"]["embed_model"],
        "retrieval": args.retrieval,
        "top_k": args.top_k,
        "temperature": args.temperature,
        "think": False,
        "modes": args.modes.split(","),
        "collection_points": health["qdrant"]["collection"].get("points_count"),
        "started_at": datetime.now().isoformat(timespec="seconds"),
    }
    (run_dir / "config.json").write_text(json.dumps(config, indent=2, ensure_ascii=False), encoding="utf-8")
    log(f"Corrida: {run_dir}  ({len(done)} respuestas ya guardadas)")
    log(f"Config: {config['model']} | {config['retrieval']} | top_k={config['top_k']} | T={config['temperature']}")

    total = len(questions) * len(config["modes"])
    n = len(done)
    with traces_path.open("a", encoding="utf-8") as fh:
        for q in questions:
            for mode in config["modes"]:
                if (q["id"], mode) in done:
                    continue
                body = {
                    "question": q["question"], "mode": mode, "top_k": args.top_k, "model": args.model,
                    "temperature": args.temperature, "retrieval": args.retrieval, "think": False,
                    "include_prompt": True,
                }
                t0 = time.perf_counter()
                try:
                    r = requests.post(f"{args.api}/query", json=body, timeout=args.timeout)
                    r.raise_for_status()
                    res = r.json()
                except Exception as exc:  # noqa: BLE001
                    res = {"question": q["question"], "mode": mode, "model": config["model"], "answer": "",
                           "contexts": [], "error": str(exc)}
                res["eval"] = {k: q.get(k) for k in ("id", "lang", "answerability", "retrieval_difficulty", "topic",
                                                      "gold_attack_ids", "gold_point_ids", "reference",
                                                      "unanswerable_part")}
                if mode == "rag":
                    res["eval"]["retrieval"] = auto_retrieval_eval(q, res.get("contexts", []))
                fh.write(json.dumps(res, ensure_ascii=False) + "\n")
                fh.flush()
                done[(q["id"], mode)] = res
                n += 1
                secs = time.perf_counter() - t0
                status = "ERROR " + res["error"][:60] if res.get("error") else f"{len(res.get('answer', ''))} chars"
                log(f"[{n:2}/{total}] {q['id']} {mode:>3} {secs:6.1f}s  {status}")

    build_workbook(questions, done, config, run_dir / "calificacion.xlsx")
    write_summary(questions, done, config, run_dir / "summary.json")
    return run_dir


# ------------------------------------------------------------------ Excel
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
AUTO_FILL = PatternFill("solid", fgColor="E2EFDA")
MANUAL_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "C2"


def build_workbook(questions: list[dict], done: dict, config: dict, out: Path) -> None:
    wb = Workbook()

    # ---- Hoja 1: Calificación (una fila por pregunta x modo)
    ws = wb.active
    ws.title = "Calificacion"
    headers = ["id", "modo", "idioma", "N0 corpus (auto)", "pregunta", "respuesta de referencia",
               "parte NO respondible", "RESPUESTA DEL SISTEMA",
               "N1 gold en top-k (auto)", "N1 rango (auto)", "N1 attack_ids recuperados (auto)",
               "N2 respuesta (0/1/2)", "N3 citas (0/1/2)", "DIAGNOSTICO", "comentario"]
    widths = [6, 6, 7, 11, 45, 55, 30, 70, 11, 9, 28, 11, 11, 16, 40]
    _header(ws, headers, widths)
    dv_score = DataValidation(type="list", formula1='"0,1,2,N/A"', allow_blank=True)
    dv_diag = DataValidation(type="list", formula1='"' + ",".join(DIAGNOSES) + '"', allow_blank=True)
    ws.add_data_validation(dv_score)
    ws.add_data_validation(dv_diag)

    row = 2
    for q in questions:
        for mode in config["modes"]:
            res = done.get((q["id"], mode))
            if res is None:
                continue
            ret = res["eval"].get("retrieval") or {}
            is_rag = mode == "rag"
            if q["answerability"] == "none":
                n1_hit, n1_rank = "N/A (nula)", "N/A"
            elif is_rag:
                n1_hit = "SI" if ret.get("gold_point_hit") else "NO"
                n1_rank = ret.get("gold_point_first_rank") or "-"
            else:
                n1_hit, n1_rank = "N/A (llm)", "N/A"
            values = [
                q["id"], mode.upper(), q["lang"], ANSWERABILITY_ES[q["answerability"]], q["question"],
                q["reference"], q.get("unanswerable_part") or "-",
                res.get("answer") or f"(ERROR: {res.get('error', '')})",
                n1_hit, n1_rank, ", ".join(str(x) for x in ret.get("retrieved_attack_ids", [])) if is_rag else "N/A",
                "", "" if is_rag else "N/A", "", "",
            ]
            ws.append(values)
            for col in range(1, len(headers) + 1):
                c = ws.cell(row=row, column=col)
                c.alignment = WRAP
                if col in (4, 9, 10, 11):
                    c.fill = AUTO_FILL
                if col in (12, 13, 14):
                    c.fill = MANUAL_FILL
            dv_score.add(ws.cell(row=row, column=12))
            dv_score.add(ws.cell(row=row, column=13))
            dv_diag.add(ws.cell(row=row, column=14))
            ws.row_dimensions[row].height = 220
            row += 1

    # ---- Hoja 2: Contextos recuperados (para verificar citas [n])
    wc = wb.create_sheet("Contextos")
    _header(wc, ["id", "modo", "[n] rango", "score", "attack_id", "nombre", "tipo", "chunk", "ES GOLD", "texto"],
            [6, 6, 8, 9, 12, 28, 16, 8, 9, 120])
    r = 2
    for q in questions:
        res = done.get((q["id"], "rag"))
        if not res:
            continue
        gold = set(q.get("gold_point_ids") or [])
        for c in res.get("contexts", []):
            m = c["metadata"]
            wc.append([q["id"], "RAG", c["rank"], round(c["score"], 4), m.get("attack_id"), m.get("name"),
                       m.get("object_type"), f"{m.get('chunk_index', 0) + 1}/{m.get('n_chunks', 1)}",
                       "SI" if c["point_id"] in gold else "", c["text"]])
            wc.cell(row=r, column=10).alignment = WRAP
            if c["point_id"] in gold:
                for col in range(1, 11):
                    wc.cell(row=r, column=col).fill = AUTO_FILL
            r += 1

    # ---- Hoja 3: Guía de calificación
    wg = wb.create_sheet("Guia")
    wg.column_dimensions["A"].width = 22
    wg.column_dimensions["B"].width = 110
    guide = [
        ("Configuración", f"modelo={config['model']} · recuperación={config['retrieval']} · top_k={config['top_k']} · "
                          f"temperatura={config['temperature']} · embeddings={config['embed_model']}"),
        ("Cómo usar", "Una fila por pregunta y modo. Las columnas verdes vienen calculadas; llena sólo las amarillas "
                      "(N2, N3, DIAGNOSTICO) y el comentario. Para verificar citas [n], busca la fila en la hoja "
                      "'Contextos' con el mismo id y rango n."),
        ("N0 corpus", "completa = toda la respuesta está en el corpus · parcial = una parte sí y otra no · "
                      "nula = nada está. Viene del diseño de la pregunta; no se modifica."),
        ("N1 recuperación", "Sólo modo RAG. SI = al menos un fragmento gold apareció entre los top-k; rango = posición "
                            "del primero. Calculado automáticamente con los gold_point_ids de la pregunta."),
        ("N2 respuesta", "2 = correcta y completa respecto a la referencia · 1 = a medias (omite algo relevante o mezcla "
                         "algo inventado con lo correcto) · 0 = incorrecta o inventada.\n"
                         "Pregunta NULA: 2 = se abstuvo ('el contexto no contiene...') · 0 = inventó una respuesta.\n"
                         "Pregunta PARCIAL: 2 = responde la parte disponible Y se abstiene de la otra · 1 = responde bien "
                         "la parte disponible pero inventa la otra · 0 = inventa o erra la parte disponible."),
        ("N3 citas", "Sólo modo RAG. 2 = cada afirmación tiene cita [n] y el fragmento n realmente la sostiene · "
                     "1 = alguna afirmación sin cita, o alguna cita apunta a un fragmento que no dice eso · "
                     "0 = sin citas o citas equivocadas. Si el sistema se abstuvo correctamente, N3 = N/A."),
        ("DIAGNOSTICO", "El PRIMER eslabón que falló, en este orden:\n"
                        "OK = sin error relevante.\n"
                        "E0-corpus = la info no estaba en el corpus y el sistema NO se abstuvo (inventó).\n"
                        "E1-recuperacion = la info estaba (completa/parcial) pero el fragmento gold no llegó (N1 = NO) "
                        "y la respuesta sufrió por ello.\n"
                        "E2-generacion = el fragmento gold llegó (N1 = SI) y aun así N2 < 2.\n"
                        "E3-citas = respuesta correcta (N2 = 2) pero N3 < 2.\n"
                        "Modo LLM: sólo puede ser OK, E0-corpus (inventó en nula/parcial) o E2-generacion."),
        ("Comparación", "RAG vs LLM: compara la columna N2 de las dos filas de la misma pregunta. Si RAG > LLM, la "
                        "recuperación aportó; si RAG < LLM, el contexto perjudicó (ruido) o el modelo se ancló mal."),
    ]
    for k, v in guide:
        wg.append([k, v])
        wg.cell(row=wg.max_row, column=1).font = Font(bold=True)
        wg.cell(row=wg.max_row, column=1).alignment = WRAP
        wg.cell(row=wg.max_row, column=2).alignment = WRAP

    # ---- Hoja 4: Resumen con fórmulas sobre la hoja Calificacion
    wr = wb.create_sheet("Resumen")
    wr.column_dimensions["A"].width = 38
    wr.column_dimensions["B"].width = 14
    wr.column_dimensions["C"].width = 14
    last = row - 1
    wr.append(["Métrica", "RAG", "LLM"])
    wr.append(["Respuestas calificadas (N2 no vacío)",
               f'=COUNTIFS(Calificacion!B2:B{last},"RAG",Calificacion!L2:L{last},"<>")',
               f'=COUNTIFS(Calificacion!B2:B{last},"LLM",Calificacion!L2:L{last},"<>")'])
    wr.append(["Promedio N2 respuesta (0-2)",
               f'=IFERROR(AVERAGEIFS(Calificacion!L2:L{last},Calificacion!B2:B{last},"RAG"),"")',
               f'=IFERROR(AVERAGEIFS(Calificacion!L2:L{last},Calificacion!B2:B{last},"LLM"),"")'])
    wr.append(["Promedio N3 citas (0-2)",
               f'=IFERROR(AVERAGEIFS(Calificacion!M2:M{last},Calificacion!B2:B{last},"RAG"),"")', "N/A"])
    wr.append(["N1 hit@k (gold en top-k, preguntas completas/parciales)",
               f'=COUNTIFS(Calificacion!B2:B{last},"RAG",Calificacion!I2:I{last},"SI")&" / "&'
               f'COUNTIFS(Calificacion!B2:B{last},"RAG",Calificacion!D2:D{last},"<>nula")', "N/A"])
    wr.append([])
    wr.append(["Diagnóstico", "RAG", "LLM"])
    for d in DIAGNOSES:
        wr.append([d, f'=COUNTIFS(Calificacion!B2:B{last},"RAG",Calificacion!N2:N{last},"{d}")',
                   f'=COUNTIFS(Calificacion!B2:B{last},"LLM",Calificacion!N2:N{last},"{d}")'])
    for cell in ("A1", "B1", "C1", "A7", "B7", "C7"):
        wr[cell].font = Font(bold=True)

    wb.save(out)
    log(f"Hoja de calificacion: {out}")


def write_summary(questions: list[dict], done: dict, config: dict, out: Path) -> None:
    rag = [done[(q["id"], "rag")] for q in questions if (q["id"], "rag") in done]
    answerable = [r for r in rag if r["eval"]["answerability"] != "none" and not r.get("error")]
    hits = sum(1 for r in answerable if r["eval"]["retrieval"]["gold_point_hit"])
    summary = {
        "config": config,
        "n_questions": len(questions),
        "n_responses": len(done),
        "n_errors": sum(1 for r in done.values() if r.get("error")),
        "retrieval_gold_hit_at_k": {"hits": hits, "total": len(answerable)},
        "avg_total_ms": {
            m: round(sum(r["timings_ms"]["total"] for r in done.values() if r["mode"] == m and "timings_ms" in r)
                     / max(1, sum(1 for r in done.values() if r["mode"] == m and "timings_ms" in r)), 1)
            for m in config["modes"]
        },
    }
    out.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    log(f"hit@{config['top_k']} (gold chunk en top-k): {hits}/{len(answerable)}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--api", default="http://localhost:8000")
    ap.add_argument("--questions", default=str(HERE / "questions.jsonl"))
    ap.add_argument("--modes", default="rag,llm")
    ap.add_argument("--top-k", type=int, default=5)
    ap.add_argument("--temperature", type=float, default=0.0)
    ap.add_argument("--retrieval", default="hybrid", choices=["hybrid", "dense", "sparse"])
    ap.add_argument("--model", default=None)
    ap.add_argument("--timeout", type=int, default=1800)
    ap.add_argument("--resume", default=None, help="carpeta de una corrida anterior para continuarla")
    args = ap.parse_args()
    run_dir = run(args)
    log(f"Listo. Resultados en {run_dir}")


if __name__ == "__main__":
    sys.exit(main())
